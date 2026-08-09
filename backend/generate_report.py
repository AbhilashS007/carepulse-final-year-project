import json
import os
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def get_chapters_from_transcript():
    path = r'C:\Users\admin\.gemini\antigravity-ide\brain\7b6c0971-bb7b-4823-9e1e-90ee70faf473\.system_generated\logs\transcript_full.jsonl'
    chapters = {1: None, 2: None, 3: None, 4: None, 5: None}
    with open(path, 'r', encoding='utf-8') as f:
        for line in f:
            try:
                data = json.loads(line)
                if data.get('source') == 'MODEL' and data.get('type') == 'PLANNER_RESPONSE':
                    content = data.get('content', '')
                    for ch in range(1, 6):
                        if f'CHAPTER {ch}' in content and f'**{ch}. ' in content:
                            chapters[ch] = content
            except Exception as e:
                pass
    return chapters

def clean_content(text):
    # Remove agent chat chatter like "Here is the finalized Chapter X"
    parts = text.split('***')
    if len(parts) > 1:
        text = parts[1] # the actual chapter usually falls between *** markers
    
    # Also strip "Source Files Used" block if it's at the top
    if 'Source Files Used' in text:
        text = re.sub(r'\*Source Files Used\*.*?\*\*\*', '', text, flags=re.DOTALL)
    
    text = text.replace('📄 ', '')
    return text.strip()

ch_data = get_chapters_from_transcript()

chapter_6 = """
### CHAPTER 6

**6. IMPLEMENTATION AND CODING**

**6.1 Introduction:**
Implementation is the phase of the project where the theoretical design is translated into functional software. This chapter details the coding practices, development environment, and core logic snippets used to build the CarePulse system. The implementation strictly adheres to modern software engineering principles, ensuring efficient integration between the backend APIs, edge processing, and frontend user interface.

**6.2 Platform and Tools Used:**
The CarePulse system was developed using a modern, decoupled technology stack to ensure scalability and rapid iteration.
* Backend Framework: FastAPI (Python 3.10+) – Chosen for its high performance, native async support, and automatic OpenAPI documentation generation.
* Frontend Library: React 18 with TypeScript – Chosen for building a robust, component-based user interface.
* Styling: Tailwind CSS – A utility-first CSS framework for rapidly building custom, responsive UI designs.
* Database: SQLite – Embedded relational database used during development, mapped via SQLAlchemy ORM for seamless future migrations to PostgreSQL or MySQL.
* Package Managers: npm (Node Package Manager) for frontend dependencies and pip for Python backend dependencies.
* Edge Hardware SDK: PlatformIO / Arduino IDE – Used to compile and flash the C++ firmware onto the ESP32 microcontroller.
* IDE: Visual Studio Code – The primary code editor used across all project layers.

**Figure 6.1 Implementation Workflow**
[Insert Diagram]
Source: Developed for the CarePulse System

**6.3 Code Implementation:**
The following sections highlight critical segments of the CarePulse backend codebase, demonstrating how complex requirements were implemented programmatically.

**6.3.1 Asynchronous Telemetry Ingestion (FastAPI)**
To prevent the ESP32 hardware from timing out while the server processes data, CarePulse utilizes FastAPI's BackgroundTasks. The API endpoint instantly validates the payload, saves the raw data, and returns a 201 Created status to the hardware. 

**Listing 6.1: Asynchronous Telemetry Ingestion**
```python
@router.post("", response_model=TelemetryOut, status_code=status.HTTP_201_CREATED)
def create_telemetry(data: TelemetryCreate, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    record = crud.create_telemetry(db, data)
    background_tasks.add_task(process_telemetry_packet, record.id)
    return record
```
Source: Developed from the CarePulse source code (backend/app/routers/telemetry.py)
Explanation: This implementation uses FastAPI's BackgroundTasks to process telemetry asynchronously, allowing the ESP32 device to receive an immediate response while the heavier analytical processing and database updates occur in the background without blocking the network thread.

**6.3.2 Composite Risk Engine Algorithm (Python)**
The core AI intelligence of CarePulse relies on a deterministic rule engine that aggregates multiple risk vectors. The algorithm combines raw telemetry (0–40 points), alert history (0–25 points), disease severity (0–20 points), and baseline disease profiles (0–15 points) to generate a 0–100 risk score.

**Listing 6.2: Composite Risk Calculation**
```python
def calculate_base_risk(patient: Patient, telemetry: Dict[str, Any], alerts: Dict[str, Any]) -> Dict[str, Any]:
    # Logic to normalize telemetry, alerts, severity, and disease baseline
    raw_score = telemetry_score + alert_score + severity_score + disease_base
    score = int(min(100, max(0, round(raw_score))))
    if score >= 80: risk_level = "critical"
    elif score >= 60: risk_level = "high"
    elif score >= 40: risk_level = "moderate"
    else: risk_level = "low"
    return {"score": score, "level": risk_level}
```
Source: Developed from the CarePulse source code (backend/app/services/disease_ai.py)
Explanation: This snippet demonstrates the strict mathematical weighting system used by the application to compute an actionable risk score. By clamping maximum values natively, the application ensures the score never exceeds its defined 0–100 boundary.

**6.3.3 REST API Implementation**
The application adheres strictly to REST principles, mapping HTTP verbs to specific CRUD actions.

Table 6.1: REST API Routes
Method | Endpoint | Purpose
POST | /auth/login | Authenticate user credentials and return a Bearer token.
GET | /patients | Retrieve paginated active and archived patients.
POST | /telemetry | Ingest and validate sensor packets from the ESP32.
GET | /alerts | Retrieve active system alerts sorted by severity and time.
GET | /devices | Retrieve hardware status and patient assignment mapping.
POST | /gemini/report/{id} | Generate a natural-language AI insight for a patient.

**6.4 Database Implementation:**
By defining models as Python classes, the application enforces referential integrity automatically. 

**Listing 6.3: Patient SQLAlchemy Model**
```python
class Patient(Base):
    __tablename__ = "patients"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    name: Mapped[str] = mapped_column(String(100), nullable=False)
    alerts: Mapped[list["Alert"]] = relationship("Alert", back_populates="patient", cascade="all, delete-orphan")
```
Source: Developed from the CarePulse source code (backend/app/models.py)
Explanation: This model configuration utilizes SQLAlchemy's ORM capabilities to explicitly define database-level types, nullability, and unique constraints. The inclusion of cascade="all, delete-orphan" prevents orphaned records.

**6.5 User Interface Implementation:**
The frontend translates the robust backend logic into a cohesive, user-friendly caregiver dashboard.

**6.5.1 Component-Based Architecture**
The application UI is broken down into small, reusable React pieces rather than monolithic files. For example, the Sidebar layout, PatientFormModal, and LiveTelemetryCard are self-contained.

**6.5.2 Responsive UI using Tailwind CSS**
Instead of writing custom CSS stylesheets, Tailwind CSS utility classes are injected directly into the JSX markup to style components conditionally.

**Listing 6.4: Conditional Tailwind CSS Rendering**
```tsx
<NavLink className={cn('sidebar-item group', isActive && 'active')}>
    <Icon className={cn('w-5 h-5 transition-colors', isActive ? 'text-white' : 'text-slate-400 group-hover:text-white')} />
</NavLink>
```
Source: Developed from the CarePulse source code (src/components/layout/Sidebar.tsx)
Explanation: This frontend snippet demonstrates how React state (isActive) is seamlessly combined with a utility function (cn) to dynamically alter the Tailwind CSS classes of the navigation links.

**6.5.3 API Integration and State Management**
To handle state and HTTP requests, the frontend implements dedicated service files (authService.ts, patientService.ts) that utilize the axios API. These services attach the JWT Bearer token to every request and parse the JSON responses back into strongly-typed TypeScript interfaces.

**6.6 Implementation Screenshots**

**Figure 6.2 Login Screen**
[Insert Screenshot of Login Page]
Source: CarePulse Application

**Figure 6.3 Caregiver Dashboard**
[Insert Screenshot of Dashboard]
Source: CarePulse Application

**Figure 6.4 Patient Management Module**
[Insert Screenshot of Patient Management]
Source: CarePulse Application

**Figure 6.5 Alert Center**
[Insert Screenshot of Alerts]
Source: CarePulse Application

**Figure 6.6 Analytics Trends**
[Insert Screenshot of Analytics]
Source: CarePulse Application

**Figure 6.7 Device Monitoring**
[Insert Screenshot of Device Monitoring]
Source: CarePulse Application

**6.7 Chapter Summary:**
This chapter highlighted the practical implementation of the CarePulse system. It outlined the modern web development technologies chosen for the project and provided specific code snippets demonstrating background concurrency, algorithmic risk calculation, and ORM database mapping. The inclusion of REST API routes, frontend React components, and final application screenshots verifies that the conceptual architectures designed in previous chapters were successfully developed into a functional software product.
"""

chapter_7 = """
### CHAPTER 7

**7. SOFTWARE TESTING**

**7.1 Introduction:**
Software testing is a critical phase of the development life cycle designed to identify errors, gaps, or missing requirements in contrast to the actual requirements. It evaluates the functionality of the software application with the intent to find whether the developed software met the specified requirements and identifies any defects to ensure that the product is defect-free. For the CarePulse system, a rigorous testing approach was adopted to ensure that real-time telemetry processing, database concurrency, and AI interactions execute flawlessly without blocking edge hardware.

**7.2 Testing Strategies:**
The testing process was divided into several levels to isolate specific components and verify their integration.
* Unit Validation: Core functions, particularly the mathematical rule engine calculations (e.g., calculate_base_risk, compute_confidence), were manually validated using representative and boundary-value inputs during development to ensure the algorithm never produces scores outside the 0–100 limits.
* Integration Testing: Integration between the React frontend and FastAPI backend was validated through functional testing of API communication, particularly verifying JWT Bearer token headers and payload structures.
* System Testing: Validated the complete end-to-end flow from the ESP32 hardware sending a POST payload to the server updating the SQLite database, and the dashboard accurately reflecting the new UrinationEvent upon refresh.
* Regression Testing: Ensures that new code modifications (such as archiving logic) do not break existing telemetry workflows via automated programmatic execution.

**7.3 Test Cases:**
The following test cases represent a subset of the critical functional tests executed during the system testing phase.

Table 7.1: Test Case TC01 – JWT Authentication
* Description: Verify that invalid credentials return an HTTP 401 Unauthorized status and valid credentials return an access token.
* Input Data: Email: admin@carepulse.com, Password: admin123
* Expected Output: JSON response containing access_token and token_type "Bearer".
* Actual Output: Token successfully generated and returned; HTTP 200 OK.
* Status: PASS

Table 7.2: Test Case TC02 – Hardware Mapping Conflict Prevention
* Description: Verify that assigning an active device_id to a newly created patient is rejected.
* Input Data: POST /patients with device_id equal to a currently assigned device.
* Expected Output: The transaction is rolled back and an HTTP 409 Conflict error is returned to the user.
* Actual Output: HTTP 409 Conflict returned with descriptive error message.
* Status: PASS

Table 7.3: Test Case TC03 – Non-Destructive Archiving (Soft Delete)
* Description: Verify that archiving a patient preserves their historical alerts and frees up their device_id for reuse.
* Input Data: PATCH /patients/{patient_id}/archive
* Expected Output: The patient's is_archived flag is set to True, the physical row remains in the database, and their historical alerts remain intact. The device_id becomes available.
* Actual Output: Historical alerts preserved; device_id successfully reused for a new patient.
* Status: PASS

Table 7.4: Test Case TC04 – Asynchronous Telemetry Ingestion
* Description: Verify that an ESP32 telemetry packet with a 90% wetness value successfully triggers a high_wetness alert in the database without blocking the HTTP response.
* Input Data: POST /telemetry containing "wetness_percent": 90
* Expected Output: Immediate HTTP 201 Created response. A background task creates an Alert record in the database.
* Actual Output: Dashboard stats updated to reflect 1 new event and 1 active alert upon refresh.
* Status: PASS

**7.4 Automated Regression Testing:**
To validate the system architecture continuously, a custom regression testing suite was implemented programmatically (backend/run_regression_test.py). 

The script systematically verifies the critical path of the application using the requests library:
1. Authentication: Authenticates dynamically and extracts the Bearer token.
2. Patient Creation: Creates a temporary test patient and binds a test device_id.
3. Telemetry Simulation: Simulates an ESP32 hardware packet that crosses the critical wetness threshold.
4. State Verification: Queries the /alerts and /dashboard/stats endpoints to verify that the background rule engine accurately generated the corresponding high_wetness alert.
5. Device Connectivity: Queries the /devices endpoint to verify the mock hardware is correctly reporting online with accurate Wi-Fi RSSI metrics.
6. Archiving and Reuse: Archives the test patient, verifies the alerts are preserved, and ensures the device_id can be safely reassigned to a new patient without conflict.

By executing this automated script, developers ensure that complex SQL relational cascades and background thread deduplication logic remain intact as the software evolves.

**7.5 Testing Summary and Validation Evidence:**
To consolidate the testing outcomes across all modules, the following table summarizes the overall functional status of the application components. Furthermore, execution screenshots validate the successful execution of the core scenarios.

Table 7.5: Consolidated Testing Summary
Test Type | Result
Authentication | Pass
Patient Management | Pass
Telemetry Ingestion | Pass
Alert Generation | Pass
Device Monitoring | Pass
AI Insights | Pass

**Figure 7.1 Automated Regression Script Execution**
[Insert Screenshot of Terminal executing run_regression_test.py]
Source: CarePulse System Testing Environment

**Figure 7.2 Successful Authentication Validation**
[Insert Screenshot of Login Success/Token Storage]
Source: CarePulse System Testing Environment

**Figure 7.3 Patient Creation Conflict Validation**
[Insert Screenshot of HTTP 409 Error on UI]
Source: CarePulse System Testing Environment

**Figure 7.4 Dashboard Telemetry Verification**
[Insert Screenshot of Dashboard updating after simulated telemetry]
Source: CarePulse System Testing Environment

**7.6 Chapter Summary:**
This chapter described the comprehensive testing methodologies applied to the CarePulse software project. By outlining detailed test cases and implementing a programmatic regression test suite, the reliability, security, and data integrity of the system were thoroughly validated. The successful execution of these test cases confirms that the system meets its foundational requirements and is resilient against unexpected inputs and concurrency edge cases.
"""

chapter_8 = """
### CHAPTER 8

**8. CONCLUSION AND FUTURE ENHANCEMENTS**

**8.1 Conclusion:**
The development of the CarePulse Smart Monitoring System successfully demonstrates the integration of Internet of Things (IoT) hardware with a modern, decoupled web architecture. By replacing manual, periodic caregiver checks with continuous, real-time sensor telemetry, the system is designed to support patient hygiene management while reducing the need for frequent manual monitoring by caregivers. 

The implementation of a deterministic disease-aware Rule Engine, layered with Google Gemini’s natural language generation, bridges the gap between raw hardware signals and meaningful health-monitoring insights. The backend architecture, built on FastAPI and SQLite, successfully utilizes asynchronous background tasks to process high-frequency data without blocking edge devices. Meanwhile, the React and Tailwind CSS frontend provides a highly responsive, component-based dashboard that centralizes patient data, analytics, and critical alerts into a single cohesive interface. Overall, the implementation demonstrates that the proposed CarePulse architecture successfully integrates IoT hardware, web technologies, database management, and AI-assisted monitoring into a functional prototype suitable for academic evaluation and future enhancement.

**8.2 Limitations:**
While the current implementation of CarePulse is robust, certain limitations exist within the scope of the project:
* Network Dependency: The ESP32 edge devices require a continuous Wi-Fi connection to transmit telemetry. In environments with poor network infrastructure, sensor data may be dropped, as local hardware caching is not fully implemented.
* External API Reliance: The generation of natural language monitoring narratives relies entirely on the external Google Gemini API. If the API rate limits are exceeded or the service experiences downtime, the system falls back to generating insights strictly using the internal Rule Engine.
* Hardware Constraints: The current iteration relies on battery-powered edge hardware. While battery levels are tracked and reported in the dashboard, there is no physical backup power supply for the sensor in the event of a sudden battery failure.
* Aggressive Polling Thresholds: The system currently flags devices as "offline" if a telemetry packet is not received within 30 seconds. In clinical environments with high network latency, this strict threshold could lead to false-positive offline alerts.
* Development Validation: The current implementation has been validated primarily in a controlled development environment using simulated and limited real telemetry rather than large-scale clinical deployment.

**8.3 Future Enhancements:**
To transition CarePulse from a functional prototype to an enterprise-grade medical monitoring solution, the following enhancements are proposed for future iterations:
* Local Hardware Caching (Edge Computing): Implementing an SD card module or utilizing the ESP32's flash memory to cache telemetry packets locally when Wi-Fi drops, pushing the data to the server in bulk once the connection is restored.
* Mobile Application Development: Developing a companion mobile application using React Native or Flutter to provide caregivers with instant push notifications, bypassing the need to actively monitor a web dashboard.
* Database Migration: Upgrading the current SQLite embedded database to a clustered PostgreSQL or MySQL environment to support high availability, horizontal scaling, and faster execution of complex time-series analytical queries.
* Predictive Machine Learning (TinyML): Integrating lightweight machine learning models directly onto the ESP32 microcontroller to perform preliminary wetness classification locally before transmitting summarized telemetry to the backend, thereby reducing required bandwidth.
* Wearable Integration: Expanding the sensor suite to monitor additional vital signs (e.g., skin temperature, heart rate) alongside moisture levels, providing a more comprehensive holistic view of patient health.

**8.4 Project Outcome:**
The CarePulse system successfully achieved the objectives established at the beginning of the project. The developed prototype demonstrates:
* Real-time ESP32 telemetry collection
* Secure FastAPI REST APIs with JWT authentication
* Patient and device lifecycle management
* Automated alert generation based on deterministic rule-based analysis
* AI-assisted monitoring insights utilizing Google Gemini
* An interactive React dashboard featuring statistical analytics and device monitoring

The project establishes a strong technical foundation for future research and development in IoT-enabled healthcare monitoring systems.
"""

doc = Document()

# Set default styles
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
paragraph_format = style.paragraph_format
paragraph_format.line_spacing = 1.15
paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Add Headings style
for i in range(1, 4):
    h_style = doc.styles[f'Heading {i}']
    h_font = h_style.font
    h_font.name = 'Times New Roman'
    h_font.size = Pt(14)
    h_font.bold = True
    h_font.color.rgb = None

def add_markdown_text(doc, text):
    lines = text.split('\\n')
    in_code_block = False
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
        
        if line.startswith('```'):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style.font.name = 'Consolas'
            p.style.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(0)
            continue
            
        if line.startswith('### '):
            doc.add_heading(line.replace('### ', '').replace('**', ''), level=1)
        elif line.startswith('## '):
            doc.add_heading(line.replace('## ', '').replace('**', ''), level=2)
        elif line.startswith('# '):
            doc.add_heading(line.replace('# ', '').replace('**', ''), level=1)
        elif line.startswith('**') and line.endswith('**') and len(line) < 100:
            doc.add_heading(line.replace('**', ''), level=2)
        else:
            # Handle inline bold
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part.replace('**', ''))
                    run.bold = True
                else:
                    p.add_run(part)

# Abstract
doc.add_heading('ABSTRACT', level=1)
doc.add_paragraph("The CarePulse Smart Monitoring System represents a modern, IoT-driven solution for enhancing patient hygiene management in healthcare environments. Built upon a decoupled architecture, the system employs ESP32 microcontrollers for real-time telemetry acquisition, transmitting data to a high-performance Python FastAPI backend. A custom deterministic Rule Engine analyzes raw sensor data in conjunction with patient clinical profiles to dynamically generate actionable alerts. These alerts, alongside comprehensive analytics, are presented through a responsive React and Tailwind CSS dashboard. Furthermore, the system integrates Google Gemini AI to distill complex telemetry trends into meaningful, natural language monitoring insights. Ultimately, this project bridges the gap between raw hardware signals and healthcare administration, demonstrating a scalable framework for continuous, non-invasive patient monitoring.")
doc.add_page_break()

# ACKNOWLEDGEMENT
doc.add_heading('ACKNOWLEDGEMENT', level=1)
doc.add_paragraph("I would like to express my sincere gratitude to my project guide and faculty for their continuous support and guidance throughout the duration of this Master of Computer Applications (MCA) project. Their insights and feedback have been invaluable in shaping the CarePulse Smart Monitoring System. I also extend my thanks to the department and the institution for providing the necessary resources and environment to complete this project successfully.")
doc.add_page_break()

# TABLE OF CONTENTS
doc.add_heading('TABLE OF CONTENTS', level=1)
doc.add_paragraph("1. INTRODUCTION\\n2. SYSTEM REQUIREMENTS SPECIFICATION\\n3. SYSTEM ANALYSIS\\n4. DATABASE DESIGN\\n5. DETAILED DESIGN\\n6. IMPLEMENTATION AND CODING\\n7. SOFTWARE TESTING\\n8. CONCLUSION AND FUTURE ENHANCEMENTS\\n9. REFERENCES")
doc.add_page_break()

# Add Chapters 1-5 from transcript
for ch in range(1, 6):
    if ch_data[ch]:
        clean_text = clean_content(ch_data[ch])
        add_markdown_text(doc, clean_text)
        doc.add_page_break()

# Add Chapters 6-8
add_markdown_text(doc, chapter_6)
doc.add_page_break()

add_markdown_text(doc, chapter_7)
doc.add_page_break()

add_markdown_text(doc, chapter_8)
doc.add_page_break()

# References
doc.add_heading('9. REFERENCES', level=1)
add_markdown_text(doc, """
* React Documentation. (n.d.). Retrieved from https://react.dev
* FastAPI Documentation. (n.d.). Retrieved from https://fastapi.tiangolo.com
* SQLAlchemy Documentation. (n.d.). Retrieved from https://www.sqlalchemy.org
* Espressif ESP32 Documentation. (n.d.). Retrieved from https://docs.espressif.com
* Google Gemini API Documentation. (n.d.). Retrieved from https://ai.google.dev
""")

doc.save('CarePulse_MCA_Project_Report.docx')
print("Document saved successfully.")
